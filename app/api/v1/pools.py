"""Kirka pools API: observer accounts (ViaBTC), daily hashrate per MSK day, workers, import."""
from __future__ import annotations

import io
from datetime import date, datetime, timedelta, timezone

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.api.v1.auth import get_current_admin
from app.db import farm_models as fm
from app.db import pool_models as pm
from app.db.session import get_db
from app.services.pools.daily import day_table
from app.services.pools.sync import import_other, import_rows, rows_from_xlsx

router = APIRouter(prefix="/pools", tags=["pools"], dependencies=[Depends(get_current_admin)])

TH = 1e12


def _th(v):
    return None if v is None else round(v / TH, 3)


def _live(a: pm.PoolAccount) -> bool:
    from app.services.connections import LIVE
    return a.provider in LIVE and a.coin not in ("", "?")


def _acc(a: pm.PoolAccount) -> dict:
    from app.services.connections import POOLS
    return {
        "id": a.id, "provider": a.provider, "pool": POOLS.get(a.provider, (a.provider,))[0], "coin": a.coin, "account": a.account, "user": a.note_user,
        "owner_user_id": a.owner_user_id, "url": a.url, "created_at": a.created_at.isoformat() if getattr(a, "created_at", None) else None,
        "group": a.group_name, "track": a.track, "active": a.active,
        "ths_10m": _th(a.hs_10m), "ths_1h": _th(a.hs_1h), "ths_1d": _th(a.hs_1d),
        "w_active": a.workers_active, "w_inactive": a.workers_inactive, "w_invalid": a.workers_invalid,
        "balance": a.balance, "profit_24h": a.profit_24h, "profit_total": a.profit_total,
        "last_ok": a.last_ok.isoformat() if a.last_ok else None, "error": a.last_error, "note": a.note,
    }


@router.get("/accounts")
def accounts(db: Session = Depends(get_db)):
    from app.services.connections import POOLS
    rows = db.scalars(select(pm.PoolAccount).order_by(pm.PoolAccount.account)).all()
    via = [_acc(a) for a in rows if _live(a)]
    other = [{"id": a.id, "provider": a.provider, "name": POOLS.get(a.provider, (a.provider,))[0], "coin": a.coin if a.coin != "?" else "",
              "user": a.note_user, "url": a.url, "owner_user_id": a.owner_user_id, "error": a.last_error}
             for a in rows if not _live(a)]
    return {"items": via, "other": other}


def _parse_period(d_from: str | None, d_to: str | None) -> tuple[date, date]:
    msk = timezone(timedelta(hours=3))
    today = datetime.now(msk).date()
    try:
        b = date.fromisoformat(d_to) if d_to else today
        a = date.fromisoformat(d_from) if d_from else b - timedelta(days=6)
    except ValueError:
        raise HTTPException(400, "Дата в формате ГГГГ-ММ-ДД")
    if a > b:
        a, b = b, a
    if (b - a).days > 366:
        raise HTTPException(400, "Не больше года за раз")
    return a, b


@router.get("/accounts/{aid}/daily")
def account_daily(aid: int, d_from: str | None = None, d_to: str | None = None, tz: int = 3, db: Session = Depends(get_db)):
    a = db.get(pm.PoolAccount, aid)
    if not a:
        raise HTTPException(404, "Нет такой ссылки")
    f, t = _parse_period(d_from, d_to)
    rows = day_table(db, aid, f, t, tz)
    for r in rows:
        r["avg_ths"] = _th(r.pop("avg_hs")); r["sum_ths"] = _th(r.pop("sum_hs")); r["pool_ths"] = _th(r.pop("pool_hs"))
    full = [r for r in rows if r["avg_ths"] is not None]
    tot = {"days": len(full), "avg_ths": round(sum(r["avg_ths"] for r in full) / len(full), 3) if full else None,
           "profit": sum(r["profit"] or 0 for r in rows), "money_usd": round(sum(r["money_usd"] or 0 for r in rows), 2)}
    hours = db.scalar(select(func.count()).select_from(pm.PoolHourly).where(pm.PoolHourly.account_id == aid))
    first = db.scalar(select(func.min(pm.PoolHourly.ts)).where(pm.PoolHourly.account_id == aid))
    return {"account": _acc(a), "rows": rows, "total": tot, "archive_hours": hours,
            "archive_since": datetime.fromtimestamp(first, tz=timezone.utc).isoformat() if first else None}


@router.get("/daily")
def all_daily(d_from: str | None = None, d_to: str | None = None, tz: int = 3, coin: str | None = None, db: Session = Depends(get_db)):
    f, t = _parse_period(d_from, d_to)
    from app.services.connections import LIVE
    q = select(pm.PoolAccount).where(pm.PoolAccount.active.is_(True), pm.PoolAccount.provider.in_(LIVE), pm.PoolAccount.coin.notin_(("", "?")))
    if coin:
        q = q.where(pm.PoolAccount.coin == coin)
    out = []
    for a in db.scalars(q.order_by(pm.PoolAccount.account)):
        rows = day_table(db, a.id, f, t, tz)
        out.append({"account": _acc(a), "days": [{"day": r["day"], "ths": _th(r["avg_hs"]), "hours": r["hours"], "part": r["source"].startswith("в архиве"), "src": r["source"], "profit": r["profit"]} for r in rows]})
    return {"from": f.isoformat(), "to": t.isoformat(), "items": out}


@router.get("/daily/export.xlsx")
def daily_export(d_from: str | None = None, d_to: str | None = None, tz: int = 3, coin: str | None = None, db: Session = Depends(get_db)):
    import openpyxl
    from openpyxl.styles import Font

    data = all_daily(d_from, d_to, tz, coin, db)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Хешрейт по суткам"
    days = [d["day"] for d in data["items"][0]["days"]] if data["items"] else []
    ws.append(["Аккаунт", "Название", "Монета"] + [datetime.fromisoformat(d).strftime("%d.%m") for d in days] + ["Среднее, TH/s"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for it in data["items"]:
        a = it["account"]
        vals = [d["ths"] for d in it["days"]]
        ok = [v for v in vals if v is not None]
        ws.append([a["account"], a["user"], a["coin"]] + vals + [round(sum(ok) / len(ok), 3) if ok else None])
    ws2 = wb.create_sheet("Начисления")
    ws2.append(["Аккаунт", "Название", "Монета"] + [datetime.fromisoformat(d).strftime("%d.%m") for d in days] + ["Итого"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for it in data["items"]:
        a = it["account"]
        vals = [d["profit"] for d in it["days"]]
        ws2.append([a["account"], a["user"], a["coin"]] + vals + [sum(v or 0 for v in vals)])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"kirka_pools_{data['from']}_{data['to']}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.get("/accounts/{aid}/workers")
def account_workers(aid: int, db: Session = Depends(get_db)):
    a = db.get(pm.PoolAccount, aid)
    if not a:
        raise HTTPException(404, "Нет такой ссылки")
    ws = db.scalars(select(pm.PoolWorker).where(pm.PoolWorker.account_id == aid).order_by(pm.PoolWorker.name)).all()
    names = [w.full_name for w in ws]
    devs = {}
    if names:
        for d in db.scalars(select(fm.FarmDevice).where(func.lower(fm.FarmDevice.worker).in_(names))):
            devs[(d.worker or "").lower()] = d
    items = []
    for w in ws:
        d = devs.get(w.full_name)
        items.append({
            "name": w.name, "full": w.full_name, "status": w.status, "ths_10m": _th(w.hs_10m), "ths_1h": _th(w.hs_1h),
            "ths_1d": _th(w.hs_1d), "reject_1h": w.reject_1h, "online_7d_h": w.online_7d_h,
            "last_active": w.last_active.isoformat() if w.last_active else None,
            "device": ({"id": d.id, "model": d.model, "ip": d.ip, "status": d.status, "ths": d.hashrate_ths} if d else None),
        })
    return {"account": _acc(a), "items": items, "matched": sum(1 for i in items if i["device"])}


@router.post("/import")
async def import_table(file: UploadFile = File(...), db: Session = Depends(get_db)):
    data = await file.read()
    if len(data) > 5_000_000:
        raise HTTPException(413, "Файл больше 5 МБ")
    try:
        rows, other = rows_from_xlsx(data)
    except Exception:
        raise HTTPException(400, "Не получилось прочитать таблицу (нужен .xlsx)")
    res = import_rows(db, rows)
    res["rows"] = len(rows)
    res["other_links"] = len(other)
    res["other_added"] = import_other(db, other)
    return res


@router.patch("/accounts/{aid}")
def account_patch(aid: int, body: dict, db: Session = Depends(get_db)):
    a = db.get(pm.PoolAccount, aid)
    if not a:
        raise HTTPException(404, "Нет такой ссылки")
    for k in ("active", "track"):
        if k in body:
            setattr(a, k, bool(body[k]))
    if "label" in body:
        a.note_user = (str(body["label"] or "").strip()[:120] or None)
    db.commit()
    return _acc(a)


@router.post("/add")
def account_add(body: dict, db: Session = Depends(get_db), user=Depends(get_current_admin)):
    """Админ добавляет ссылку наблюдателя прямо на странице «Пулы» (общая, видна всем админам)."""
    from app.services import connections as cn

    try:
        r = cn.pool_connect(db, user, str(body.get("url") or ""), shared=True, label=body.get("label"))
    except cn.ConnError as e:
        raise HTTPException(400, e.text) from None
    return r


@router.delete("/accounts/{aid}")
def account_delete(aid: int, db: Session = Depends(get_db)):
    a = db.get(pm.PoolAccount, aid)
    if not a:
        raise HTTPException(404, "Нет такой ссылки")
    db.delete(a)
    db.commit()
    return {"ok": True}
