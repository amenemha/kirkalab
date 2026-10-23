"""Kirka farm API: today panel, devices, ASIC card, errors, export.

Access: admins (owner Telegram IDs) for now. Roles per client come later.
"""
from __future__ import annotations

import io
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.api.v1.auth import get_current_admin
from app.db import farm_models as fm
from app.db.session import get_db
from app.services.farm.classes import CLASSES

router = APIRouter(prefix="/farm", tags=["farm"], dependencies=[Depends(get_current_admin)])


def _refs(db: Session):
    sites = {s.id: s.name.strip() for s in db.scalars(select(fm.FarmSite))}
    conts = {c.id: (c.name.strip(), c.site_id) for c in db.scalars(select(fm.FarmContainer))}
    clients = {c.id: c.name.strip() for c in db.scalars(select(fm.FarmClient))}
    return sites, conts, clients


def _dev_row(d: fm.FarmDevice, sites, conts, clients) -> dict:
    return {
        "id": d.id, "mac": d.mac, "serial": d.serial, "num": d.user_number, "model": d.model,
        "site": sites.get(d.site_id), "site_id": d.site_id,
        "container": (conts.get(d.container_id) or (None,))[0], "container_id": d.container_id,
        "client": clients.get(d.client_id), "client_id": d.client_id,
        "status": d.status, "ip": d.ip, "worker": d.worker, "pool": d.pool_url,
        "mode": d.power_mode, "ths": d.hashrate_ths, "ths5m": d.hashrate_5m_ths, "nominal": d.nominal_ths,
        "w": d.power_w, "temp": d.temp_env, "uptime": d.uptime_sec, "alarms": d.alarms or [],
        "level": d.worst_level, "tags": [t.get("label") for t in (d.tags or []) if t.get("label")],
        "since": d.status_since.isoformat() if d.status_since else None,
    }


@router.get("/summary")
def summary(db: Session = Depends(get_db)):
    sites, conts, clients = _refs(db)
    devs = list(db.scalars(select(fm.FarmDevice)))
    last = db.scalar(select(fm.FarmSyncRun).where(fm.FarmSyncRun.ok.is_(True)).order_by(fm.FarmSyncRun.id.desc()).limit(1))
    tot = Counter(d.status for d in devs)
    hr = sum(d.hashrate_ths or 0 for d in devs if d.status == "working")
    nom = sum(d.nominal_ths or 0 for d in devs if d.status == "working")
    kw = sum((d.power_w or 0) for d in devs) / 1000
    by_cont = defaultdict(lambda: dict(total=0, working=0, offline=0, suspended=0, ths=0.0, kw=0.0, alarms=0, temp=None))
    by_client = defaultdict(lambda: dict(total=0, working=0, offline=0, ths=0.0))
    cls = Counter()
    for d in devs:
        c = by_cont[d.container_id]
        c["total"] += 1
        c[d.status if d.status in ("working", "offline", "suspended") else "offline"] += 1
        c["ths"] += d.hashrate_ths or 0
        c["kw"] += (d.power_w or 0) / 1000
        if d.alarm_count:
            c["alarms"] += 1
        if d.temp_env is not None:
            c["temp"] = max(c["temp"] or -99, d.temp_env)
        k = by_client[d.client_id]
        k["total"] += 1
        k["working"] += d.status == "working"
        k["offline"] += d.status == "offline"
        k["ths"] += d.hashrate_ths or 0
        for a in d.alarms or []:
            cls[a.get("cls") or "other"] += 1
    containers = []
    for cid, v in by_cont.items():
        name, sid = conts.get(cid, ("Без контейнера", None))
        containers.append(dict(id=cid, name=name, site=sites.get(sid), site_id=sid, **{k: (round(x, 2) if isinstance(x, float) else x) for k, x in v.items()}))
    containers.sort(key=lambda x: (x["site"] or "", x["name"]))
    cl = [dict(id=k, name=clients.get(k, "Без клиента"), **{a: (round(b, 2) if isinstance(b, float) else b) for a, b in v.items()}) for k, v in by_client.items()]
    cl.sort(key=lambda x: -x["total"])
    incidents = [dict(id=i.id, scope=i.scope_name, started=i.started_at.isoformat(), ended=i.ended_at.isoformat() if i.ended_at else None, total=i.devices_total, down=i.devices_down)
                 for i in db.scalars(select(fm.FarmIncident).order_by(fm.FarmIncident.started_at.desc()).limit(20))]
    return {
        "synced_at": last.finished_at.isoformat() if last else None,
        "total": len(devs), "working": tot.get("working", 0), "offline": tot.get("offline", 0),
        "suspended": tot.get("suspended", 0), "ths": round(hr, 1), "nominal_ths": round(nom, 1),
        "efficiency": round(hr / nom * 100, 1) if nom else None, "kw": round(kw, 1),
        "with_alarms": sum(1 for d in devs if d.alarm_count),
        "classes": [dict(key=k, label=CLASSES[k][0], count=cls.get(k, 0)) for k in sorted(CLASSES, key=lambda x: CLASSES[x][1])],
        "sites": [dict(id=k, name=v) for k, v in sites.items()],
        "containers": containers, "clients": cl, "incidents": incidents,
    }


def _filtered(db: Session, q, status, site, container, client, model, cls, only_alarms):
    st = select(fm.FarmDevice)
    if status:
        st = st.where(fm.FarmDevice.status.in_(status.split(",")))
    if site:
        st = st.where(fm.FarmDevice.site_id.in_([int(x) for x in site.split(",")]))
    if container:
        st = st.where(fm.FarmDevice.container_id.in_([int(x) for x in container.split(",")]))
    if client:
        st = st.where(fm.FarmDevice.client_id.in_([int(x) for x in client.split(",")]))
    if model:
        st = st.where(fm.FarmDevice.model.in_(model.split(",")))
    if only_alarms:
        st = st.where(fm.FarmDevice.alarm_count > 0)
    if q:
        like = f"%{q.strip()}%"
        st = st.where(or_(fm.FarmDevice.ip.ilike(like), fm.FarmDevice.mac.ilike(like), fm.FarmDevice.serial.ilike(like),
                          fm.FarmDevice.worker.ilike(like), fm.FarmDevice.user_number.ilike(like), fm.FarmDevice.model.ilike(like)))
    devs = list(db.scalars(st))
    if cls:
        want = set(cls.split(","))
        devs = [d for d in devs if any((a.get("cls") in want) for a in (d.alarms or []))]
    return devs


@router.get("/devices")
def devices(
    q: str | None = None, status: str | None = None, site: str | None = None, container: str | None = None,
    client: str | None = None, model: str | None = None, cls: str | None = None, alarms: bool = False,
    sort: str = "container", limit: int = Query(200, le=6000), offset: int = 0, db: Session = Depends(get_db),
):
    sites, conts, clients = _refs(db)
    devs = _filtered(db, q, status, site, container, client, model, cls, alarms)
    rows = [_dev_row(d, sites, conts, clients) for d in devs]
    desc = sort.startswith("-")
    key = sort.lstrip("-")
    rows.sort(key=lambda r: (r.get(key) is None, r.get(key) if r.get(key) is not None else 0) if key in ("ths", "temp", "w", "uptime", "level") else (str(r.get(key) or ""), str(r.get("ip") or "")), reverse=desc)
    facets = {
        "models": Counter(r["model"] for r in rows if r["model"]).most_common(),
        "status": Counter(r["status"] for r in rows).most_common(),
    }
    return {"total": len(rows), "items": rows[offset:offset + limit], "facets": facets}


@router.get("/devices/{dev_id}")
def device_card(dev_id: int, db: Session = Depends(get_db)):
    d = db.get(fm.FarmDevice, dev_id)
    if d is None:
        raise HTTPException(404, "not found")
    sites, conts, clients = _refs(db)
    ev = db.scalars(select(fm.FarmDeviceEvent).where(fm.FarmDeviceEvent.device_id == dev_id).order_by(fm.FarmDeviceEvent.at.desc()).limit(300))
    since = datetime.now(timezone.utc) - timedelta(hours=72)
    hourly = db.scalars(select(fm.FarmDeviceHourly).where(fm.FarmDeviceHourly.device_id == dev_id, fm.FarmDeviceHourly.hour >= since).order_by(fm.FarmDeviceHourly.hour))
    row = _dev_row(d, sites, conts, clients)
    row.update({
        "first_seen": d.first_seen_at.isoformat() if d.first_seen_at else None,
        "last_seen": d.last_seen_at.isoformat() if d.last_seen_at else None,
        "online_total": d.online_sec_total, "pools": d.pools or [], "tags_full": d.tags or [],
        "events": [dict(at=e.at.isoformat(), kind=e.kind, old=e.old, new=e.new, code=e.code, cls=e.cls, level=e.level, author=e.author) for e in ev],
        "hourly": [dict(h=h.hour.isoformat(), ths=round(h.hashrate_sum / h.samples, 2) if h.samples else 0, up=round(h.online_samples / h.samples * 100) if h.samples else 0, t=h.temp_max) for h in hourly],
    })
    return row


@router.get("/errors")
def errors(
    q: str | None = None, site: str | None = None, container: str | None = None, client: str | None = None,
    model: str | None = None, cls: str | None = None, level: int | None = None, group: str = "code",
    db: Session = Depends(get_db),
):
    sites, conts, clients = _refs(db)
    devs = _filtered(db, q, None, site, container, client, model, None, True)
    want = set(cls.split(",")) if cls else None
    rows = []
    for d in devs:
        base = _dev_row(d, sites, conts, clients)
        for a in d.alarms or []:
            if want and a.get("cls") not in want:
                continue
            if level is not None and (a.get("level") or 0) < level:
                continue
            rows.append({**{k: base[k] for k in ("id", "model", "site", "container", "client", "ip", "worker", "serial", "status", "temp")},
                         "code": a.get("code"), "msg": a.get("msg"), "cls": a.get("cls"), "cls_label": CLASSES.get(a.get("cls") or "other", ("Прочее",))[0],
                         "lvl": a.get("level"), "since": a.get("since")})
    groups = defaultdict(lambda: dict(count=0, devices=set()))
    gkey = {"code": lambda r: f"{r['code']} · {r['msg']}", "cls": lambda r: r["cls_label"], "container": lambda r: r["container"] or "—",
            "client": lambda r: r["client"] or "Без клиента", "model": lambda r: r["model"] or "—"}.get(group, lambda r: r["code"])
    for r in rows:
        g = groups[gkey(r)]
        g["count"] += 1
        g["devices"].add(r["id"])
    grp = sorted(({"key": k, "count": v["count"], "devices": len(v["devices"])} for k, v in groups.items()), key=lambda x: -x["count"])
    rows.sort(key=lambda r: (-(r["lvl"] or 0), -(r["since"] or 0)))
    return {"total": len(rows), "devices": len({r["id"] for r in rows}), "groups": grp, "items": rows}


@router.get("/errors/export.xlsx")
def errors_export(site: str | None = None, container: str | None = None, client: str | None = None,
                  model: str | None = None, cls: str | None = None, ids: str | None = None, db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    data = errors(None, site, container, client, model, cls, None, "code", db)
    items = data["items"]
    if ids:
        keep = {int(x) for x in ids.split(",") if x}
        items = [r for r in items if r["id"] in keep]
    wb = Workbook()
    ws = wb.active
    ws.title = "Ошибки"
    head = ["Площадка", "Контейнер", "Клиент", "Модель", "IP", "Воркер", "Серийный №", "Статус", "Класс", "Код", "Ошибка", "Уровень", "С какого времени (МСК)", "Темп., °C"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="01696F")
    msk = timezone(timedelta(hours=3))
    for r in items:
        since = datetime.fromtimestamp(r["since"], msk).strftime("%d.%m.%Y %H:%M") if r["since"] else ""
        ws.append([r["site"], r["container"], r["client"], r["model"], r["ip"], r["worker"], r["serial"], r["status"], r["cls_label"], r["code"], r["msg"], r["lvl"], since, r["temp"]])
    for col, w in zip("ABCDEFGHIJKLMN", (18, 34, 24, 20, 15, 30, 34, 10, 14, 8, 36, 8, 18, 8)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"kirka_errors_{datetime.now(msk).strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.get("/containers/{cid}/history")
def container_history(cid: int, hours: int = 24, db: Session = Depends(get_db)):
    since = datetime.now(timezone.utc) - timedelta(hours=hours)
    rows = db.scalars(select(fm.FarmContainerSample).where(fm.FarmContainerSample.container_id == cid, fm.FarmContainerSample.at >= since).order_by(fm.FarmContainerSample.at))
    return [dict(at=r.at.isoformat(), online=r.online, offline=r.offline, total=r.total, ths=r.hashrate_ths, kw=r.power_kw) for r in rows]


@router.get("/history")
def farm_history(hours: int = 24, db: Session = Depends(get_db)):
    since = datetime.now(timezone.utc) - timedelta(hours=hours)
    q = (select(fm.FarmContainerSample.at, func.sum(fm.FarmContainerSample.hashrate_ths), func.sum(fm.FarmContainerSample.online),
                func.sum(fm.FarmContainerSample.total), func.sum(fm.FarmContainerSample.power_kw))
         .where(fm.FarmContainerSample.at >= since).group_by(fm.FarmContainerSample.at).order_by(fm.FarmContainerSample.at))
    return [dict(at=a.isoformat(), ths=round(h or 0, 1), online=int(o or 0), total=int(t or 0), kw=round(k or 0, 1)) for a, h, o, t, k in db.execute(q)]
