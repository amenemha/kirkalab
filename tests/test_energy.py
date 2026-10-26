import io

import openpyxl

from app.services.energy.hourly_import import parse, read_rows
from app.services.energy.plugs import PLUGS, plug_for, seed_plugs
from tests.test_rbac_users import _make_admin

HEAD = [
    ["Дата", "Подстанция\nФидер\nСчетчик", "ГПП-16", "", "", "Итого"],
    ["", "", "Фидер-6", "Фидер-75 10кВ", "Итого", ""],
    ["", "", 1117240906, 1113250116, "", ""],
    ["", "", 1, 1, "", ""],
    ["", "", "0", "0", "", ""],
    ["", "", "0", "0", "", ""],
    ["", "", "A+", "A+", "", ""],
]


def _xlsx(days=2, broken=False) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in HEAD:
        ws.append(r)
    ws.append([""] * 6)
    ws2 = wb.create_sheet("Page 1")
    n = 0
    for d in range(1, days + 1):
        for h in range(24):
            row = [f"202609{d:02d}", f"{h}:00 - {h + 1}:00", 100 + h, "7,5", 0, ""]
            (ws if n < 30 else ws2).append(row)
            n += 1
    ws2.append(["Итого", "", 0, 0, 0, ""])
    if broken:
        ws["A1"] = "что-то"
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_parse_xlsx_multisheet_msk():
    p = parse(read_rows(_xlsx(), "a.xlsx"))
    assert [m["serial"] for m in p["meters"]] == ["1117240906", "1113250116"]
    assert p["meters"][0]["feeder"] == "Фидер-6" and p["meters"][0]["substation"] == "ГПП-16"
    s = p["data"]["1117240906"]
    assert len(s) == 48
    first = min(s)
    assert first.isoformat() == "2026-08-31T21:00:00+00:00" and s[first] == 100  # 0:00 МСК
    assert p["data"]["1113250116"][first] == 7.5


def test_plugs_lookup():
    assert plug_for("Antminer S21+ Hyd.", PLUGS)["inlet"].startswith("3ф")
    assert plug_for("Antminer T21", PLUGS)["inlet"] == "P34"
    assert plug_for("Whatsminer M30S++", PLUGS)["cord"] == "C19"
    assert plug_for("Elphapex DG1+", PLUGS)["inlets"] == 2
    assert plug_for("Antminer S19k Pro", PLUGS)["inlet"] == "C14"
    assert plug_for("IceRiver KS5L", PLUGS)["inlet"] == "C20"
    assert plug_for(None, PLUGS) is None


def test_import_api_and_hourly(client, db):
    token, _ = _make_admin(client, db)
    h = {"Authorization": f"Bearer {token}"}
    up = lambda data, name="p.xlsx": client.post("/api/v1/energy/import", headers=h,  # noqa: E731
                                                 files={"file": (name, data)})
    r = up(_xlsx())
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["meters"] == 2 and j["rows"] == 96 and j["gaps"] == 0 and j["start"] == "2026-09-01 00:00"
    assert j["end"] == "2026-09-02 23:00"
    # повторная загрузка не дублирует
    assert up(_xlsx()).json()["rows"] == 96
    ms = client.get("/api/v1/energy/meters", headers=h).json()
    assert len(ms) == 2 and ms[0]["hours"] == 48
    d = client.get("/api/v1/energy/hourly?start=2026-09-01&end=2026-09-01&group=day", headers=h).json()
    assert len(d["series"]) == 1 and d["total_kwh"] == sum(100 + i for i in range(24)) + 7.5 * 24
    hh = client.get("/api/v1/energy/hourly?start=2026-09-02&end=2026-09-02", headers=h).json()
    assert len(hh["series"]) == 24 and hh["series"][0]["t"] == "2026-09-02 00:00"
    assert hh["peak"]["t"] == "2026-09-02 23:00"
    assert client.get("/api/v1/energy/hourly?start=2026-09-05&end=2026-09-01", headers=h).status_code == 422
    assert up(_xlsx(broken=True)).status_code == 422
    assert up(b"", "x.xls").status_code == 422
    assert client.get("/api/v1/energy/meters").status_code in (401, 403)


def test_readings_nodes_plugs(client, db):
    token, _ = _make_admin(client, db)
    h = {"Authorization": f"Bearer {token}"}
    m = client.post("/api/v1/energy/meters", json={"serial": "0113", "name": "3а ввод 2", "ratio": 1}, headers=h).json()
    assert client.post("/api/v1/energy/meters", json={"serial": "0113"}, headers=h).status_code == 409
    r1 = client.post("/api/v1/energy/readings", json={"meter_id": m["id"], "value": "100,5",
                                                         "at": "2026-10-25T10:00:00"}, headers=h)
    assert r1.status_code == 200
    r2 = client.post("/api/v1/energy/readings", json={"meter_id": m["id"], "value": 340.5,
                                                         "at": "2026-10-26T10:00:00", "source": "photo"}, headers=h).json()
    assert r2["delta_kwh"] == 240 and r2["hours"] == 24 and r2["avg_kw"] == 10
    bad = client.post("/api/v1/energy/readings", json={"meter_id": m["id"], "value": 50,
                                                          "at": "2026-10-26T12:00:00"}, headers=h)
    assert bad.status_code == 422
    assert len(client.get(f"/api/v1/energy/readings?meter={m['id']}", headers=h).json()) == 2
    assert client.put(f"/api/v1/energy/meters/{m['id']}", json={"group": "КТП 2"}, headers=h).json()["group"] == "КТП 2"
    assert client.put(f"/api/v1/energy/meters/{m['id']}", json={"serial": "x"}, headers=h).status_code == 422
    ktp = client.post("/api/v1/energy/nodes", json={"kind": "ktp", "name": "КТП 2", "voltage_kv": 10}, headers=h).json()
    inp = client.post("/api/v1/energy/nodes", json={"kind": "input", "name": "Ввод 2", "parent_id": ktp["id"]}, headers=h).json()
    client.post("/api/v1/energy/nodes", json={"kind": "socket", "name": "PDU2 #7", "phase": "C", "parent_id": inp["id"]}, headers=h)
    assert client.post("/api/v1/energy/nodes", json={"kind": "socket", "name": "x", "phase": "Q"}, headers=h).status_code == 422
    assert client.post("/api/v1/energy/nodes", json={"kind": "bad", "name": "x"}, headers=h).status_code == 422
    assert len(client.get("/api/v1/energy/nodes", headers=h).json()) == 3
    assert client.delete(f"/api/v1/energy/nodes/{ktp['id']}", headers=h).json()["deleted"] == 3
    engine = db.get_bind()
    assert seed_plugs(engine) == len(PLUGS) and seed_plugs(engine) == 0
    p = client.get("/api/v1/energy/plugs", headers=h).json()
    assert len(p["reference"]) == len(PLUGS)
