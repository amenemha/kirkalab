import io

import openpyxl

from app.db import farm_models as fm
from tests.test_rbac_users import _make_admin


def _seed(db):
    c = fm.FarmContainer(ext_id="c1", name="5а 8-9 ↑ Ф1 (КТП 3) P33 IBMM")
    db.add(c)
    db.flush()
    devs = []
    for i in range(1, 11):
        d = fm.FarmDevice(ext_id=f"d{i}", model="Antminer T21", container_id=c.id, ip=f"10.116.8.{i}",
                          status="working" if i % 3 else "offline", hashrate_ths=190.0, nominal_ths=190.0,
                          alarms=[{"code": "E1", "msg": "Fan fail", "cls": "fan", "level": 2, "since": 1760000000}] if i == 2 else [])
        db.add(d)
        devs.append(d)
    db.commit()
    return c, devs


def test_export_exactly_selected(client, db):
    token, _ = _make_admin(client, db)
    h = {"Authorization": f"Bearer {token}"}
    _, devs = _seed(db)
    pick = [devs[1].id, devs[4].id, devs[8].id]  # selected 3 after unselecting others
    r = client.post("/api/v1/farm/devices/export.xlsx", json={"ids": pick}, headers=h)
    assert r.status_code == 200 and r.content[:2] == b"PK"
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 3
    ips = sorted(x[9] for x in rows)
    assert ips == sorted(f"10.116.8.{i}" for i in (2, 5, 9))
    assert {x[4] for x in rows} <= {"Работает", "Не работает"}
    assert rows[0][2].startswith("стеллаж 1 · полка 1 · место 2")
    err = list(wb.worksheets[1].iter_rows(min_row=2, values_only=True))
    assert len(err) == 1 and err[0][9] == "Опасно"
    assert client.post("/api/v1/farm/devices/export.xlsx", json={"ids": []}, headers=h).status_code == 400
    assert client.post("/api/v1/farm/devices/export.xlsx", json={"ids": pick}).status_code in (401, 403)
